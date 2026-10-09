import os, shutil, subprocess, sys, threading, tempfile
from pathlib import Path
from datetime import datetime
from flask import Flask, jsonify, request, send_from_directory, Response
import openpyxl

ROOT = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv('DATA_DIR', '/var/data'))
DATA_DIR.mkdir(parents=True, exist_ok=True)
CS_NAME = 'Updated CS Identification all.xlsx'
GAS_NAME = 'AMNS_Pune_Gas_Hazardous_Area_Safety_Register.xlsx'
CS_STORE = DATA_DIR / CS_NAME
GAS_STORE = DATA_DIR / GAS_NAME
app = Flask(__name__, static_folder=None)
app.config['MAX_CONTENT_LENGTH'] = 30 * 1024 * 1024
lock = threading.RLock()
last_signature = None

# Seed persistent storage once, from the package's included workbooks.
for name, store in [(CS_NAME, CS_STORE), (GAS_NAME, GAS_STORE)]:
    bundled = ROOT / name
    if not store.exists() and bundled.exists():
        shutil.copy2(bundled, store)

def validate_workbook(path, kind):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    expected = 'Confined Space Location List' if kind == 'cs' else 'Hazard Assessment Matrix'
    if expected not in wb.sheetnames:
        raise ValueError(f"Workbook must contain the sheet '{expected}'.")
    ws = wb[expected]
    if kind == 'cs':
        count = sum(1 for r in range(9, ws.max_row + 1) if ws.cell(r, 4).value not in (None, ''))
    else:
        headers = {str(ws.cell(4,c).value).strip(): c for c in range(1, ws.max_column+1) if ws.cell(4,c).value}
        if 'Area ID' not in headers:
            raise ValueError("Gas workbook is missing the 'Area ID' column.")
        count = sum(1 for r in range(5, ws.max_row+1) if ws.cell(r, headers['Area ID']).value not in (None, ''))
    wb.close()
    if count < 1:
        raise ValueError('No data records were found in the workbook.')
    return count

def refresh_data(force=False):
    global last_signature
    with lock:
        if not CS_STORE.exists() or not GAS_STORE.exists():
            raise RuntimeError('Source Excel files are missing from persistent storage.')
        signature = (CS_STORE.stat().st_mtime_ns, CS_STORE.stat().st_size, GAS_STORE.stat().st_mtime_ns, GAS_STORE.stat().st_size)
        if not force and signature == last_signature and (ROOT/'data.json').exists() and (ROOT/'gas_data.json').exists():
            return
        # Keep the generator scripts and Excel in the same folder they expect.
        shutil.copy2(CS_STORE, ROOT / CS_NAME)
        shutil.copy2(GAS_STORE, ROOT / GAS_NAME)
        subprocess.run([sys.executable, str(ROOT/'tools'/'generate_data.py')], cwd=ROOT, check=True, capture_output=True, text=True)
        subprocess.run([sys.executable, str(ROOT/'tools'/'generate_gas_data.py')], cwd=ROOT, check=True, capture_output=True, text=True)
        last_signature = signature

@app.get('/')
def home():
    return send_from_directory(ROOT, 'index.html')

@app.get('/api/health')
def health():
    try:
        refresh_data()
        import json
        cs=json.loads((ROOT/'data.json').read_text(encoding='utf-8'))
        gas=json.loads((ROOT/'gas_data.json').read_text(encoding='utf-8'))
        return jsonify(status='ok', confinedSpaceRecords=len(cs.get('locations',[])), gasHazardAreas=len(gas.get('areas',[])), updatedAt=cs.get('generatedAt'), storage='persistent' if str(DATA_DIR).startswith('/var/data') else str(DATA_DIR))
    except Exception as e:
        return jsonify(status='error', error=str(e)), 500

@app.get('/api/data')
def cs_data():
    try:
        refresh_data()
        return send_from_directory(ROOT, 'data.json', max_age=0)
    except Exception as e:
        return jsonify(error=str(e)), 500

@app.get('/api/gas-data')
def gas_data():
    try:
        refresh_data()
        return send_from_directory(ROOT, 'gas_data.json', max_age=0)
    except Exception as e:
        return jsonify(error=str(e)), 500

@app.get('/<path:path>')
def static_files(path):
    blocked = {CS_NAME, GAS_NAME, 'app.py', 'requirements.txt', 'render.yaml', 'README.txt', 'README_RENDER.txt'}
    if path.startswith(('api/', 'tools/', '.github/')) or path in blocked or path.lower().endswith(('.xlsx', '.xlsm', '.xls')):
        return jsonify(error='Not found'), 404
    return send_from_directory(ROOT, path)

@app.route('/admin', methods=['GET','POST'])
def admin():
    if request.method == 'GET':
        return Response('''<!doctype html><html><head><meta name="viewport" content="width=device-width,initial-scale=1"><title>Dashboard data update</title><style>body{font:16px Arial;max-width:720px;margin:40px auto;padding:20px;color:#222}h1{color:#a71920}form{border:1px solid #ddd;border-radius:12px;padding:20px;margin:20px 0}label{display:block;font-weight:bold;margin:14px 0 6px}input{width:100%;box-sizing:border-box;padding:12px}button{background:#a71920;color:white;padding:12px 18px;border:0;border-radius:6px;margin-top:18px}small{color:#555}</style></head><body><h1>AM/NS Safety Register — Excel Update</h1><p>Upload a valid Excel register to refresh the dashboard data immediately.</p><form method="post" enctype="multipart/form-data"><label>Admin password</label><input name="password" type="password" required autocomplete="current-password"><label>Confined Space Excel (.xlsx)</label><input name="cs_file" type="file" accept=".xlsx"><label>Gas Hazardous Area Excel (.xlsx)</label><input name="gas_file" type="file" accept=".xlsx"><small>Choose at least one workbook. Keep the required sheet names and column layout unchanged.</small><br><button type="submit">Validate, Save & Refresh Dashboard</button></form><p><a href="/api/health">Check dashboard data status</a> · <a href="/">Open dashboard</a></p></body></html>''', mimetype='text/html')
    expected = os.getenv('ADMIN_PASSWORD')
    if not expected:
        return Response('Admin upload is disabled. Set ADMIN_PASSWORD in Render environment variables.', status=503)
    if request.form.get('password','') != expected:
        return Response('Incorrect admin password.', status=401)
    uploads = [('cs_file', CS_STORE, 'cs'), ('gas_file', GAS_STORE, 'gas')]
    selected = [(field, store, kind) for field,store,kind in uploads if request.files.get(field) and request.files[field].filename]
    if not selected:
        return Response('Please select at least one Excel workbook.', status=400)
    backups=[]
    try:
        with lock:
            for field, store, kind in selected:
                file=request.files[field]
                if not file.filename.lower().endswith('.xlsx'):
                    raise ValueError('Only .xlsx files are accepted.')
                fd,tmp=tempfile.mkstemp(suffix='.xlsx', dir=str(DATA_DIR)); os.close(fd)
                temp=Path(tmp)
                file.save(temp)
                count=validate_workbook(temp,kind)
                backups.append((store, store.read_bytes() if store.exists() else None))
                os.replace(temp,store)
            refresh_data(force=True)
        return Response('<h2>Success — workbook saved and dashboard refreshed.</h2><p>Counts: '+ ' · '.join([f'{kind}: {validate_workbook(store,kind)} records' for _,store,kind in selected]) + '</p><p><a href="/">Open dashboard</a> · <a href="/api/health">Verify API</a> · <a href="/admin">Upload another file</a></p>', mimetype='text/html')
    except Exception as e:
        with lock:
            for store, old in backups:
                if old is not None: store.write_bytes(old)
        try: refresh_data(force=True)
        except Exception: pass
        return Response('Update failed; previous valid data was restored. Details: '+str(e), status=400, mimetype='text/plain')

# Build datasets on startup so first request is ready.
try:
    refresh_data(force=True)
except Exception as exc:
    print('Initial data refresh warning:', exc, flush=True)

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.getenv('PORT', '10000')))
