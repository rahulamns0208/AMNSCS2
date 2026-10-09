import json, datetime
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
XLSX=ROOT/"AMNS_Pune_Gas_Hazardous_Area_Safety_Register.xlsx"
OUT=ROOT/"gas_data.json"

def clean(v):
    if v is None: return ""
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return str(v).strip()

wb=openpyxl.load_workbook(XLSX,data_only=True)
ws=wb["Hazard Assessment Matrix"]

# Header row is 4 in Excel (1-based); data starts on row 5.
headers=[clean(ws.cell(4,c).value) for c in range(1,ws.max_column+1)]
idx={h:i+1 for i,h in enumerate(headers) if h}
areas=[]
for r in range(5,ws.max_row+1):
    area_id=clean(ws.cell(r,idx.get("Area ID",1)).value)
    if not area_id: continue
    areas.append({
      "id":area_id,
      "location":clean(ws.cell(r,idx["Hazard Area Location"]).value),
      "gas":clean(ws.cell(r,idx["Primary Gas Present"]).value),
      "flammable":clean(ws.cell(r,idx["Flammable? (Yes/No)"]).value),
      "toxic":clean(ws.cell(r,idx["Toxic / Asphyxiant?"]).value),
      "zone":clean(ws.cell(r,idx["Hazard Zone / Area Class"]).value),
      "limits":clean(ws.cell(r,idx["LEL / UEL / Exposure Limit"]).value),
      "ignition":clean(ws.cell(r,idx["Potential Ignition Sources"]).value),
      "detection":clean(ws.cell(r,idx["Gas Detection System"]).value),
      "interlocks":clean(ws.cell(r,idx["Safety Interlocks & Trips"]).value),
      "ventilation":clean(ws.cell(r,idx["Ventilation Systems"]).value),
      "ppe":clean(ws.cell(r,idx["Mandatory PPE Requirements"]).value),
      "sop":clean(ws.cell(r,idx["Applicable SOPs & Permits"]).value),
      "risk":clean(ws.cell(r,idx["Risk Level"]).value)
    })

refws=wb["Safety Reference Guide"]
reference=[]
# The guide has two related tables. Preserve both in a common normalized shape.
for r in range(6,refws.max_row+1):
    a=clean(refws.cell(r,1).value); b=clean(refws.cell(r,2).value)
    c=clean(refws.cell(r,3).value); d=clean(refws.cell(r,4).value)
    if any([a,b,c,d]):
        reference.append({
          "Parameter Name":a,
          "Safety Category":b,
          "Description & Operational Importance":c,
          "Key Reference Standards":d
        })

data={
 "generatedAt":datetime.datetime.now().strftime("%d %b %Y %H:%M"),
 "sourceWorkbook":XLSX.name,
 "sheet":"Hazard Assessment Matrix",
 "areas":areas,
 "reference":reference
}
OUT.write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding="utf-8")
print(f"Generated {OUT} with {len(areas)} gas hazardous areas.")
